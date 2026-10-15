import math
import os
import zipfile
from datetime import date, datetime, time

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

import config


def _parse_hour(value):
    if isinstance(value, datetime):
        return value.hour
    if isinstance(value, time):
        return value.hour
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if 0 <= value < 1:
            return round(value * 24) % 24
        return None
    if not isinstance(value, str):
        return None

    text = value.strip()
    for fmt in ("%H:%M", "%I:%M %p", "%I %p"):
        try:
            return datetime.strptime(text, fmt).hour
        except ValueError:
            continue
    return None


def _parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in (
            "%Y-%m-%d",
            "%d-%b-%y",
            "%d-%b-%Y",
            "%m/%d/%Y",
            "%d/%m/%Y",
        ):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
    return None


def _find_date_and_hour_cells(sheet, target_date, target_hour):
    for row in sheet.iter_rows():
        date_columns = [
            cell.column
            for cell in row
            if isinstance(cell.value, str) and cell.value.strip().casefold() == "date"
        ]
        if not date_columns:
            continue

        hour_columns = {
            cell.column
            for cell in row
            if _parse_hour(cell.value) == target_hour
        }
        if not hour_columns:
            continue

        for date_column in date_columns:
            hour_column = next(
                (column for column in sorted(hour_columns) if column > date_column),
                None,
            )
            if hour_column is None:
                continue
            for row_number in range(row[0].row + 1, sheet.max_row + 1):
                if _parse_date(sheet.cell(row_number, date_column).value) == target_date:
                    return sheet.cell(row_number, hour_column)
    return None


def _parse_record(record):
    timestamp = record.get("timestamp")
    try:
        recorded_at = datetime.strptime(str(timestamp), "%Y-%m-%d %H:%M:%S")
    except (TypeError, ValueError):
        return None, f"Invalid record timestamp for Excel export: {timestamp!r}"

    sheet_name = str(
        record.get("host_name")
        or record.get("lpar")
        or record.get("server")
        or record.get("config_key")
        or ""
    ).strip()
    if not sheet_name:
        return None, "Record has no server name for Excel worksheet selection"

    values = {}
    for metric in ("asp", "cpu"):
        raw_value = record.get(metric)
        try:
            percentage = float(raw_value)
            if not math.isfinite(percentage) or percentage < 0 or (
                metric == "asp" and percentage > 100
            ):
                raise ValueError(f"Invalid {metric.upper()} percentage: {raw_value!r}")
            values[metric] = percentage
        except (TypeError, ValueError) as exc:
            return None, str(exc)

    return {
        "sheet_name": sheet_name,
        "recorded_at": recorded_at,
        "values": values,
    }, None


def _update_workbook_batch(path, metric, records):
    if not os.path.isfile(path):
        raise FileNotFoundError(f"Workbook not found: {path}")

    lock_path = os.path.join(os.path.dirname(path), f"~${os.path.basename(path)}")
    if os.path.exists(lock_path):
        raise PermissionError(f"Workbook is open in Excel: {path}")

    workbook = load_workbook(path)
    try:
        for record in records:
            sheet_name = record["sheet_name"]
            sheet = next(
                (
                    item
                    for item in workbook.worksheets
                    if item.title.casefold() == sheet_name.casefold()
                ),
                None,
            )
            if sheet is None:
                raise KeyError(
                    f"Worksheet {sheet_name!r} was not found in {path}"
                )

            recorded_at = record["recorded_at"]
            cell = _find_date_and_hour_cells(
                sheet,
                recorded_at.date(),
                recorded_at.hour,
            )
            if cell is None:
                raise ValueError(
                    f"No cell found for {recorded_at.date().isoformat()} at "
                    f"{recorded_at.hour:02d}:00 in worksheet {sheet_name!r}"
                )

            cell.value = record["values"][metric] / 100
            if cell.number_format == "General":
                cell.number_format = "0.00%"

        workbook.save(path)
    finally:
        workbook.close()


def push_metrics_to_excel(records):
    """Update all server cells for an hour, saving each shared workbook once."""
    if isinstance(records, dict):
        records = [records]

    parsed_records = []
    errors = []
    for record in records:
        parsed, error = _parse_record(record)
        if error:
            errors.append(f"Excel export failed: {error}")
        else:
            parsed_records.append(parsed)

    if errors:
        return errors
    if not parsed_records:
        return errors

    years = {record["recorded_at"].strftime("%Y") for record in parsed_records}
    hours = {record["recorded_at"].strftime("%Y-%m-%d %H") for record in parsed_records}
    if len(years) != 1 or len(hours) != 1:
        return errors + ["Excel batch must contain readings from one hour"]
    year = years.pop()

    year_directory = os.path.join(config.ONEDRIVE_SHAREPOINT_PATH, year)
    for metric in ("asp", "cpu"):
        workbook_path = os.path.join(
            year_directory,
            f"BY - {metric.upper()} {year}.xlsx",
        )
        try:
            _update_workbook_batch(workbook_path, metric, parsed_records)
        except (
            KeyError,
            OSError,
            TypeError,
            ValueError,
            zipfile.BadZipFile,
            InvalidFileException,
        ) as exc:
            errors.append(f"{metric.upper()} Excel batch update failed: {exc}")

    return errors
